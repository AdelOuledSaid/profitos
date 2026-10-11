from profitos.runtime import *
from profitos.accounting import AccountingError, create_entry
import csv, hashlib, io, json, re
from datetime import date, datetime


def _norm(s):
    return str(s or '').strip()


def _amount(v):
    if v is None or v == '': return 0.0
    if isinstance(v, (int,float)): return round(float(v),2)
    return round(float(str(v).replace('\u202f','').replace(' ','').replace(',','.')),2)


# Comptes de paie (PCG) : racine -> libellé, pour ouvrir automatiquement un compte de
# paie manquant au plan (ex. 645300 « retraite » absent d'un plan personnalisé).
PAYROLL_ACCOUNT_ROOTS = {
    '421': 'Personnel - rémunérations dues', '425': 'Personnel - avances et acomptes',
    '427': 'Personnel - oppositions', '428': 'Personnel - charges à payer',
    '431': 'Sécurité sociale', '437': 'Autres organismes sociaux',
    '442': "Prélèvement à la source (impôt sur le revenu)", '447': 'Autres impôts, taxes et versements assimilés',
    '631': 'Impôts, taxes et versements assimilés sur rémunérations', '633': 'Impôts, taxes et versements assimilés sur rémunérations (autres organismes)',
    '641': 'Rémunérations du personnel', '644': "Rémunération du travail de l'exploitant",
    '645': 'Charges de sécurité sociale et de prévoyance', '646': "Cotisations sociales personnelles de l'exploitant",
    '647': 'Autres charges sociales', '648': 'Autres charges de personnel',
}


def normalize_account_code(code):
    """Compte du fichier de paie -> compte à 6 chiffres du plan : « 6411 » -> 641100,
    « 64110000 » -> 641100, « 421 » -> 421000. Les codes non numériques sont gardés."""
    c = re.sub(r'\s', '', str(code or ''))
    if not c.isdigit():
        return c
    if len(c) < 6:
        return c.ljust(6, '0')
    if len(c) > 6 and set(c[6:]) <= {'0'}:
        return c[:6]
    return c


def _ensure_payroll_accounts(c, lines, entity_id):
    """Ouvre au plan les comptes de paie standard manquants (classes 4 et 6)."""
    for ln in lines:
        code = ln['account_code']
        if c.execute('SELECT 1 FROM accounting_chart_of_accounts WHERE code=? AND (entity_id IS NULL OR entity_id IS ?)',
                     (code, entity_id)).fetchone():
            continue
        label = PAYROLL_ACCOUNT_ROOTS.get(code[:3])
        if not label or len(code) != 6:
            continue
        c.execute('INSERT OR IGNORE INTO accounting_chart_of_accounts(code,label,account_class,is_collective,is_active,created_at) VALUES(?,?,?,0,1,?)',
                  (code, label, int(code[0]), datetime.utcnow().isoformat()))


def _parse_rows(raw, filename):
    ext=(filename.rsplit('.',1)[-1] if '.' in filename else '').lower()
    rows=[]
    if ext == 'csv':
        text=raw.decode('utf-8-sig')
        sample=text[:4096]
        try: dialect=csv.Sniffer().sniff(sample, delimiters=';,\t|')
        except csv.Error: dialect=csv.excel
        rows=list(csv.DictReader(io.StringIO(text), dialect=dialect))
    elif ext in ('xlsx','xlsm'):
        import zipfile
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_SPREADSHEET_UNCOMPRESSED:
                raise ValueError('Fichier Excel trop volumineux une fois décompressé.')
        from openpyxl import load_workbook
        ws=load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
        vals=list(ws.iter_rows(values_only=True))
        if not vals: return []
        headers=[_norm(x) for x in vals[0]]
        rows=[dict(zip(headers,r)) for r in vals[1:] if any(x is not None for x in r)]
    else:
        raise ValueError('Format non pris en charge : utilisez CSV ou XLSX.')
    aliases={
      'account_code':['CompteNum','Compte','account_code','account','compte'],
      'label':['EcritureLib','Libelle','Libellé','label','libelle'],
      'debit':['Debit','Débit','debit'], 'credit':['Credit','Crédit','credit'],
      'auxiliary_name':['CompAuxLib','Auxiliaire','auxiliary_name']}
    out=[]
    for r in rows:
        def get(k):
            for a in aliases[k]:
                if a in r and r[a] not in (None,''): return r[a]
            return ''
        acc=normalize_account_code(_norm(get('account_code')))
        if not acc: continue
        debit=_amount(get('debit')); credit=_amount(get('credit'))
        if debit < 0 or credit < 0:
            raise ValueError('Les montants débit/crédit doivent être positifs ou nuls.')
        if debit and credit:
            raise ValueError(f'Le compte {acc} ne peut pas avoir un débit et un crédit sur la même ligne.')
        if not debit and not credit:
            continue
        out.append({'account_code':acc,'label':_norm(get('label')) or 'Paie',
                    'auxiliary_name':_norm(get('auxiliary_name')) or None,
                    'debit':debit,'credit':credit})
    return out


def register(app):
    @app.route('/comptabilite/paie', methods=['GET','POST'])
    @login_required
    def payroll_imports_list():
        from profitos.entities import current_entity_id, user_can_access_entity
        eid=current_entity_id(); c=cx(); error=None
        if not user_can_access_entity(c, session.get('user_id'), eid):
            c.close(); abort(403)
        if request.method=='POST':
            provider=_norm(request.form.get('provider')).upper()
            period=_norm(request.form.get('period_label'))
            entry_date=_norm(request.form.get('entry_date')) or date.today().isoformat()
            f=request.files.get('file')
            if provider not in ('SILAE','PAYFIT') or not period or not f or not f.filename:
                error='Prestataire, période et fichier CSV/XLSX sont obligatoires.'
            else:
                try:
                    entry_date=date.fromisoformat(entry_date).isoformat()
                except (TypeError, ValueError):
                    error="La date comptable de l'import de paie est invalide."
            if not error:
                raw=f.read(); digest=hashlib.sha256(raw).hexdigest()
                exists=c.execute('SELECT id FROM payroll_imports WHERE entity_id IS ? AND file_sha256=?',(eid,digest)).fetchone()
                period_exists=c.execute(
                    'SELECT id FROM payroll_imports WHERE entity_id IS ? AND provider=? AND period_label=? AND status IN (\'pending\',\'posted\')',
                    (eid,provider,period),
                ).fetchone()
                if exists:
                    error='Ce fichier de paie a déjà été importé pour cette entité.'
                elif period_exists:
                    error='Un import de paie existe déjà pour ce prestataire et cette période sur cette entité.'
                else:
                    try:
                        lines=_parse_rows(raw,f.filename)
                        if len(lines)<2: raise ValueError('Le fichier doit contenir au moins deux lignes comptables exploitables.')
                        td=round(sum(x['debit'] for x in lines),2); tc=round(sum(x['credit'] for x in lines),2)
                        status='pending' if abs(td-tc)<=0.01 else 'invalid'
                        msg=None if status=='pending' else f'Écriture déséquilibrée : débit {fr_number(td,2)} € / crédit {fr_number(tc,2)} €.'
                        c.execute('INSERT INTO payroll_imports(entity_id,provider,period_label,entry_date,original_filename,file_sha256,status,total_debit,total_credit,lines_json,error_message,created_by,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)',
                                  (eid,provider,period,entry_date,f.filename,digest,status,td,tc,json.dumps(lines,ensure_ascii=False),msg,current_user()['email'],datetime.utcnow().isoformat()))
                        c.commit(); flash('Import de paie préparé. Vérifiez-le avant validation.')
                        c.close(); return redirect(url_for('payroll_imports_list'))
                    except Exception as exc: error=str(exc)
        rows=c.execute('SELECT * FROM payroll_imports WHERE entity_id IS ? ORDER BY id DESC',(eid,)).fetchall(); c.close()
        return render_template('payroll_imports.html',rows=rows,error=error,today=date.today().isoformat())

    @app.route('/comptabilite/paie/<int:import_id>/valider', methods=['POST'])
    @login_required
    def payroll_import_validate(import_id):
        from profitos.entities import current_entity_id, user_can_access_entity
        eid=current_entity_id(); c=cx()
        if not user_can_access_entity(c, session.get('user_id'), eid):
            c.close(); abort(403)
        row=c.execute("SELECT * FROM payroll_imports WHERE id=? AND entity_id IS ?",(import_id,eid)).fetchone()
        if not row or row['status']!='pending':
            c.close(); abort(404)
        try:
            if not c.execute("SELECT code FROM accounting_journals WHERE code='PA'").fetchone():
                c.execute("INSERT INTO accounting_journals(code,label,journal_type,is_default,created_at) VALUES('PA','Paie','OD',0,?)",(datetime.utcnow().isoformat(),))
            lines=json.loads(row['lines_json'])
            for ln in lines:
                ln['account_code']=normalize_account_code(ln['account_code'])
            _ensure_payroll_accounts(c, lines, eid)
            entry_id=create_entry(c,'PA',row['entry_date'],f"Paie {row['provider']} — {row['period_label']}",lines,
                                  source_type='payroll_import',source_id=row['id'],created_by=current_user()['email'],entity_id=eid,commit=False)
            posted = c.execute("UPDATE payroll_imports SET status='posted',accounting_entry_id=?,validated_by=?,validated_at=? WHERE id=? AND entity_id IS ? AND status='pending'",
                               (entry_id,current_user()['email'],datetime.utcnow().isoformat(),import_id,eid))
            if posted.rowcount != 1:
                raise AccountingError("Cet import de paie n'est plus en attente de validation.")
            c.commit()
            log_activity('PAYROLL_IMPORT_POSTED',f"Import paie #{import_id} comptabilisé dans PA")
            flash('OD de paie validée et comptabilisée dans le journal PA.')
        except (AccountingError, ValueError, TypeError, json.JSONDecodeError) as exc:
            c.rollback()
            flash(f'Validation impossible : {exc}')
        finally: c.close()
        return redirect(url_for('payroll_imports_list'))
